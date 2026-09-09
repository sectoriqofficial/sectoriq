"""
Lap Time Analyzer - Basic Skeleton
-----------------------------------
Goal: Read lap data, find your best lap, and show where you're
losing time in each sector compared to that best lap.

This is intentionally simple. No fancy libraries, no AI, just
reading data, doing subtraction, and printing results clearly.
"""

import csv
import json
import os
from datetime import datetime


# -----------------------------
# STEP 1: Read the raw CSV file
# -----------------------------
def read_lap_data(filename):
    """
    Opens a CSV file and turns each row into a dictionary.
    Expected columns: lap, sector1, sector2, sector3, total_time
    (all times in seconds, as decimals, e.g. 32.451)
    """
    laps = []
    with open(filename, newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            lap = {
                "lap": int(row["lap"]),
                "sector1": float(row["sector1"]),
                "sector2": float(row["sector2"]),
                "sector3": float(row["sector3"]),
                "total_time": float(row["total_time"]),
            }
            laps.append(lap)
    return laps


# -----------------------------------------------------
# STEP 1b: Read real Assetto Corsa data (race_out.json)
# -----------------------------------------------------
def read_lap_data_from_ac_json(filename, session_index=None):
    """
    Reads Assetto Corsa's race_out.json and converts it into the
    same lap dictionary format used elsewhere in this script.

    Key differences from our CSV format, handled here:
    - AC stores times in milliseconds, we convert to seconds.
    - The "time" field is unreliable for hotlap sessions, it's
      often -1 unless that lap set a new personal record. So
      instead we treat a lap as valid using "cuts" == 0 (clean,
      no track limit violations), and we calculate the actual
      lap time ourselves by adding up the three sector times.
    - Sectors come as a plain list [s1, s2, s3], not named fields.
    - A race_out.json can contain MULTIPLE sessions (e.g. separate
      practice and hotlap runs). By default we go through every
      session and gather laps from all of them, so nothing gets
      missed. Pass a specific session_index if you only want one.
    """
    with open(filename, "r") as f:
        data = json.load(f)

    all_sessions = data["sessions"]

    # If a specific session was requested, only look at that one.
    # Otherwise, gather laps from every session in the file.
    if session_index is not None:
        sessions_to_read = [all_sessions[session_index]]
    else:
        sessions_to_read = all_sessions

    laps = []
    for session in sessions_to_read:
        for raw in session["laps"]:
            # Skip laps with track limit violations (corner cuts).
            if raw["cuts"] != 0:
                continue

            sectors = raw["sectors"]  # [sector1_ms, sector2_ms, sector3_ms]

            # Skip incomplete laps, e.g. if the session was exited
            # or reset mid-lap, not all 3 sectors get recorded.
            if len(sectors) < 3:
                continue

            # Calculate total time ourselves from the sectors, since
            # the "time" field is only filled in for record laps.
            total_time_ms = sum(sectors)

            lap = {
                "lap": raw["lap"],
                "sector1": sectors[0] / 1000,
                "sector2": sectors[1] / 1000,
                "sector3": sectors[2] / 1000,
                "total_time": total_time_ms / 1000,
            }
            laps.append(lap)

    return laps


# -----------------------------------
# STEP 2: Find the best (fastest) lap
# -----------------------------------
def find_best_lap(laps):
    """
    Returns the lap dictionary with the lowest total_time.
    """
    return min(laps, key=lambda lap: lap["total_time"])


# -----------------------------------------------
# STEP 3: Compare every lap's sectors to the best
# -----------------------------------------------
def compare_to_best(laps, best_lap):
    """
    For every lap, calculate how much slower (or faster) each
    sector was compared to the best lap's matching sector.
    Positive number = lost time. Negative number = gained time.
    """
    comparisons = []
    for lap in laps:
        diff = {
            "lap": lap["lap"],
            "sector1_diff": round(lap["sector1"] - best_lap["sector1"], 3),
            "sector2_diff": round(lap["sector2"] - best_lap["sector2"], 3),
            "sector3_diff": round(lap["sector3"] - best_lap["sector3"], 3),
            "total_diff": round(lap["total_time"] - best_lap["total_time"], 3),
        }
        comparisons.append(diff)
    return comparisons


# -----------------------------
# STEP 4: Print it out clearly
# -----------------------------
def print_report(comparisons, best_lap):
    print(f"\nBest Lap: Lap {best_lap['lap']} — {best_lap['total_time']}s\n")
    print(f"{'Lap':<5}{'Sec1 +/-':<12}{'Sec2 +/-':<12}{'Sec3 +/-':<12}{'Total +/-':<10}")
    print("-" * 51)
    for c in comparisons:
        print(
            f"{c['lap']:<5}{c['sector1_diff']:<12}{c['sector2_diff']:<12}"
            f"{c['sector3_diff']:<12}{c['total_diff']:<10}"
        )


# -----------------------------------------------
# STEP 5: Save this session into permanent history
# -----------------------------------------------
def save_session_to_history(history_path, track, car, laps, best_lap):
    """
    Appends this session's results to history.json, so past
    sessions are never lost. Each entry records when it happened,
    which track/car, the best lap, and every valid lap that
    session, so future runs can compare across time.

    If history.json doesn't exist yet, it's created automatically.

    Skips saving if this exact session (same track, car, and laps)
    was already the most recently saved entry, so clicking
    "Analyze" repeatedly on the same race_out.json doesn't create
    duplicate history entries.
    """
    # Load existing history, or start a fresh list if none exists yet.
    if os.path.exists(history_path):
        with open(history_path, "r") as f:
            history = json.load(f)
    else:
        history = []

    # Duplicate check: compare against the most recent entry only,
    # since that's the one that would match a re-analyze of the
    # same session.
    if history:
        last_entry = history[-1]
        is_same_session = (
            last_entry["track"] == track
            and last_entry["car"] == car
            and last_entry["laps"] == laps
        )
        if is_same_session:
            print(f"\nSession already saved (no changes since last analysis, skipping duplicate).")
            return

    session_entry = {
        "date": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "track": track,
        "car": car,
        "best_lap_time": best_lap["total_time"],
        "laps": laps,
    }

    history.append(session_entry)

    with open(history_path, "w") as f:
        json.dump(history, f, indent=2)

    print(f"\nSession saved to history ({track}, best lap {best_lap['total_time']}s)")


# -----------------------------------------------
# STEP 6: Delete a session from history
# -----------------------------------------------
def delete_session_from_history(history_path, index_to_delete):
    """
    Removes one session entry from history.json by its position
    in the list (0 = oldest entry saved).
    """
    with open(history_path, "r") as f:
        history = json.load(f)

    if 0 <= index_to_delete < len(history):
        removed = history.pop(index_to_delete)
        with open(history_path, "w") as f:
            json.dump(history, f, indent=2)
        print(f"\nDeleted session: {removed['track']} ({removed['date']})")


# -----------------------------------------------
# STEP 7: Export all history to an Excel file
# -----------------------------------------------
def export_history_to_excel(history_path, output_path):
    """
    Exports your full session history to a single Excel file with
    two sheets:
      - "Lap Times": every individual lap from every session, flat
        and easy to filter/sort in Excel yourself.
      - "Trends": a pivoted table (one column per track+car combo)
        with a native Excel line chart plotted directly from it,
        so it's a real, editable chart, not just a picture.
    """
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font

    if not os.path.exists(history_path):
        raise FileNotFoundError("No history.json found yet, analyze a session first.")

    with open(history_path, "r") as f:
        history = json.load(f)

    if not history:
        raise ValueError("History is empty, nothing to export yet.")

    wb = Workbook()

    # ---------- Sheet 1: Lap Times ----------
    laps_sheet = wb.active
    laps_sheet.title = "Lap Times"

    headers = ["Date", "Track", "Car", "Lap", "Sector 1", "Sector 2", "Sector 3", "Total Time"]
    laps_sheet.append(headers)
    for cell in laps_sheet[1]:
        cell.font = Font(name="Arial", bold=True)

    for session in history:
        for lap in session["laps"]:
            laps_sheet.append([
                session["date"],
                session["track"],
                session["car"],
                lap["lap"],
                lap["sector1"],
                lap["sector2"],
                lap["sector3"],
                lap["total_time"],
            ])

    for col in laps_sheet.columns:
        laps_sheet.column_dimensions[col[0].column_letter].width = 14

    # ---------- Sheet 2: Trends ----------
    trends_sheet = wb.create_sheet("Trends")

    # Group best lap times by (track, car), in chronological order.
    groups = {}
    for session in history:
        key = f"{session['track']} — {session['car']}"
        groups.setdefault(key, []).append(session["best_lap_time"])

    group_names = list(groups.keys())
    max_sessions = max(len(v) for v in groups.values())

    # Header row: Session #, then one column per track+car combo.
    trends_sheet.append(["Session #"] + group_names)
    for cell in trends_sheet[1]:
        cell.font = Font(name="Arial", bold=True)

    for i in range(max_sessions):
        row = [i + 1]
        for name in group_names:
            values = groups[name]
            row.append(values[i] if i < len(values) else None)
        trends_sheet.append(row)

    for col in trends_sheet.columns:
        trends_sheet.column_dimensions[col[0].column_letter].width = 22

    # Native Excel line chart, built from the table above, so it
    # stays a real editable chart when opened in Excel.
    chart = LineChart()
    chart.title = "Best Lap Time Trends"
    chart.y_axis.title = "Best lap time (seconds)"
    chart.x_axis.title = "Session number"

    data = Reference(
        trends_sheet, min_col=2, max_col=1 + len(group_names),
        min_row=1, max_row=1 + max_sessions
    )
    categories = Reference(trends_sheet, min_col=1, min_row=2, max_row=1 + max_sessions)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.width = 24
    chart.height = 14

    trends_sheet.add_chart(chart, "B" + str(max_sessions + 4))

    wb.save(output_path)
    print(f"\nExported to Excel: {output_path}")


# -----------------------------
# Run everything
# -----------------------------
if __name__ == "__main__":
    # This builds the path based on where THIS script file lives,
    # so it works no matter what folder you run it from.
    script_folder = os.path.dirname(os.path.abspath(__file__))
    history_path = os.path.join(script_folder, "history.json")

    # Read directly from where Assetto Corsa actually saves its
    # session data, so you never have to manually copy the file
    # over anymore. This is the real, live file the game updates
    # every time you finish a session.
    documents_folder = os.path.join(os.environ["USERPROFILE"], "Documents")
    json_path = os.path.join(documents_folder, "Assetto Corsa", "out", "race_out.json")

    # Pull track and car name straight from the AC json, so we
    # don't have to type it in manually every time.
    with open(json_path, "r") as f:
        raw_data = json.load(f)
    track_name = raw_data["track"]
    car_name = raw_data["players"][0]["car"]

    laps = read_lap_data_from_ac_json(json_path)

    if not laps:
        print("\nNo valid laps found in this session (all laps were cut or incomplete).")
        print("Complete at least one clean lap in Assetto Corsa, then run this again.")
    else:
        best_lap = find_best_lap(laps)
        comparisons = compare_to_best(laps, best_lap)
        print_report(comparisons, best_lap)

        save_session_to_history(history_path, track_name, car_name, laps, best_lap)